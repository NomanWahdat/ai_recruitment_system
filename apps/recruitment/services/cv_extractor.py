from pathlib import Path
from docx import Document
from PyPDF2 import PdfReader
from apps.recruitment.utils.text import clean_extracted_text


class CVExtractionError(Exception):
    pass


class CVTextExtractor:
    def extract(self, file_path):
        ext = Path(file_path).suffix.lower()
        text = (
            self.extract_pdf(file_path)
            if ext == ".pdf"
            else self.extract_docx(file_path)
            if ext == ".docx"
            else ""
        )
        if not text.strip():
            raise CVExtractionError(
                "No readable text. Scanned/image-only documents need OCR; upload a text-based PDF or DOCX."
            )
        return text[:200000]

    def extract_pdf(self, file_path):
        try:
            reader = PdfReader(file_path)
            if reader.is_encrypted or len(reader.pages) > 100:
                raise CVExtractionError("Encrypted PDFs or PDFs over 100 pages are not supported.")
            return clean_extracted_text(
                "\n\n".join(page.extract_text() or "" for page in reader.pages)
            )
        except CVExtractionError:
            raise
        except Exception as exc:
            raise CVExtractionError("Unable to read PDF file.") from exc

    def extract_docx(self, file_path):
        try:
            document = Document(file_path)
            lines = [p.text for p in document.paragraphs]
            lines.extend(
                cell.text for table in document.tables for row in table.rows for cell in row.cells
            )
            return clean_extracted_text("\n".join(lines))
        except Exception as exc:
            raise CVExtractionError("Unable to read DOCX file.") from exc
