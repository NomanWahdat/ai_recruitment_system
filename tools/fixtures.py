"""Generate real, synthetic document containers without external services."""

from io import BytesIO
from docx import Document


def docx_bytes(text, table_text=None):
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    if table_text:
        document.add_table(rows=1, cols=1).cell(0, 0).text = table_text
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_bytes(text):
    def escape(line):
        return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    stream = (
        "BT /F1 11 Tf 50 780 Td "
        + " ".join("(" + escape(line) + ") Tj 0 -16 Td" for line in text.splitlines())
        + " ET"
    )
    stream = stream.encode("ascii", "replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    data = b"%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(data)
    data += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets[1:]:
        data += f"{offset:010d} 00000 n \n".encode()
    return (
        data
        + f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    )
